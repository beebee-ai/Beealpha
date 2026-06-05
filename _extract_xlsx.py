import sys
try:
    from openpyxl import load_workbook
except ImportError:
    print("NO_OPENPYXL")
    sys.exit(2)

wb = load_workbook(r'E:\飞书文件\No.3 Qishi×BEE Beta Infinity Camp Schedule.xlsx', data_only=True)
for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    print(f'=== Sheet: {sheet_name} (rows={ws.max_row}, cols={ws.max_column}) ===')
    for row in ws.iter_rows(values_only=True):
        if any(c is not None and str(c).strip() != "" for c in row):
            print(" | ".join("" if c is None else str(c) for c in row))
    print()
